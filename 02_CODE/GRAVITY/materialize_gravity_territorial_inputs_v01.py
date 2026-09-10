"""
E0-A — MATERIALIZZAZIONE INPUT TERRITORIALI GRAVITY v0

PURPOSE
-------
Creare l'artefatto derivato degli input territoriali della Gravity LIGHT v0,
materializzando P_i e A_j a partire dal workbook RAW autorizzato.

INPUT
-----
Workbook:
    Gravity_v0_territorial_inputs_raw.xlsx

Sheet:
    MASTER_RAW

Campi richiesti:
    PRO_COM
    COMUNE
    PARCO_AUTO_RAW
    TURISMO_RAW
    GDO_RAW

OUTPUT
------
Workbook:
    Gravity_v0_territorial_inputs_derived_v01.xlsx

Sheets:
    TERRITORIAL_INPUTS
    QA

ASSUMPTIONS
-----------
1. La normalizzazione è esclusivamente L1, cioè quota sul totale regionale.
2. P_i = PARCO_AUTO_RAW_i / SUM(PARCO_AUTO_RAW).
3. TURISMO_NORM_i = TURISMO_RAW_i / SUM(TURISMO_RAW).
4. GDO_NORM_i = GDO_RAW_i / SUM(GDO_RAW).
5. A_j = 0.5 * TURISMO_NORM_j + 0.5 * GDO_NORM_j.
6. A_j NON viene rinormalizzato dopo il mix.
7. Gli zeri vengono preservati.
8. Non vengono applicati min-max, z-score, log, epsilon o altre trasformazioni.
9. Tolleranza floating point = 1e-12.

FILES WRITTEN
-------------
Solo:
    Gravity_v0_territorial_inputs_derived_v01.xlsx

FILES NEVER MODIFIED
--------------------
Gravity_v0_territorial_inputs_raw.xlsx
e qualsiasi altro artefatto sorgente.

NOTE
----
Lo script è strettamente deterministico e non esegue alcuna
calibrazione Gravity.
"""

from __future__ import annotations

import argparse
import math
import sys
from collections import Counter
from numbers import Real
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_SHEET = "MASTER_RAW"
OUTPUT_FILENAME = "Gravity_v0_territorial_inputs_derived_v01.xlsx"

REQUIRED_COLUMNS = [
    "PRO_COM",
    "COMUNE",
    "PARCO_AUTO_RAW",
    "TURISMO_RAW",
    "GDO_RAW",
]

NUMERIC_RAW_COLUMNS = [
    "PARCO_AUTO_RAW",
    "TURISMO_RAW",
    "GDO_RAW",
]

OUTPUT_COLUMNS = [
    "PRO_COM",
    "COMUNE",
    "PARCO_AUTO_RAW",
    "TURISMO_RAW",
    "GDO_RAW",
    "P_i",
    "TURISMO_NORM",
    "GDO_NORM",
    "A_j",
]

EXPECTED_ROWS = 215
TOL = 1e-12


def parse_args():
    parser = argparse.ArgumentParser(
        description="Materializza P_i e A_j per Gravity LIGHT v0."
    )

    parser.add_argument(
        "--input",
        required=True,
        help="Path di Gravity_v0_territorial_inputs_raw.xlsx",
    )

    parser.add_argument(
        "--output",
        default=None,
        help=(
            "Path output opzionale. "
            "Se omesso viene creato accanto al RAW con nome canonico."
        ),
    )

    return parser.parse_args()


def is_missing(value):
    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


def numeric_value(value):
    """
    Accetta soltanto valori numerici realmente presenti nel workbook.

    Non prova a interpretare stringhe, virgole decimali o altri formati:
    una conversione implicita sarebbe una trasformazione non autorizzata.
    """

    if isinstance(value, bool):
        return None

    if isinstance(value, Real):
        number = float(value)

        if math.isfinite(number):
            return number

    return None


def fmt_number(value):
    if isinstance(value, int):
        return f"{value}"

    if isinstance(value, float):
        return f"{value:.12g}"

    return str(value)


def status_exact(value, expected):
    return "PASS" if value == expected else "FAIL"


def status_close(value, expected=1.0):
    return (
        "PASS"
        if math.isclose(
            value,
            expected,
            rel_tol=0.0,
            abs_tol=TOL,
        )
        else "FAIL"
    )


def print_rank(title, records, field, reverse):
    print()
    print(title)
    print("-" * 86)
    print(
        f"{'RANK':>4}  "
        f"{'PRO_COM':>8}  "
        f"{'COMUNE':<35}  "
        f"{field:>18}"
    )
    print("-" * 86)

    ranked = sorted(
        records,
        key=lambda r: (
            r[field],
            str(r["COMUNE"]).casefold(),
        ),
        reverse=reverse,
    )[:10]

    for rank, record in enumerate(ranked, 1):
        comune = str(record["COMUNE"])[:35]

        print(
            f"{rank:>4}  "
            f"{str(record['PRO_COM']):>8}  "
            f"{comune:<35}  "
            f"{record[field]:>18.12f}"
        )


def style_header(ws, last_column):
    header = ws.cell(row=1, column=1)
    del header  # solo per rendere esplicito che lo styling è sul range seguente

    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in ws[1]:
        if cell.column > last_column:
            break

        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def create_output_workbook(
    output_path,
    records,
    qa_rows,
):
    if output_path.exists():
        raise FileExistsError(
            f"Output già esistente; overwrite vietato:\n{output_path}"
        )

    wb = Workbook()

    # ==============================================================
    # TERRITORIAL_INPUTS
    # ==============================================================

    ws = wb.active
    ws.title = "TERRITORIAL_INPUTS"

    ws.append(OUTPUT_COLUMNS)

    for record in records:
        ws.append(
            [
                record["PRO_COM"],
                record["COMUNE"],
                record["PARCO_AUTO_RAW"],
                record["TURISMO_RAW"],
                record["GDO_RAW"],
                record["P_i"],
                record["TURISMO_NORM"],
                record["GDO_NORM"],
                record["A_j"],
            ]
        )

    style_header(ws, len(OUTPUT_COLUMNS))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    widths = {
        "A": 12,
        "B": 32,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 18,
        "H": 18,
        "I": 18,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(
        min_row=2,
        min_col=6,
        max_col=9,
    ):
        for cell in row:
            cell.number_format = "0.000000000000"

    # ==============================================================
    # QA
    # ==============================================================

    qa = wb.create_sheet("QA")

    qa.append(
        [
            "METRIC",
            "VALUE",
            "EXPECTED",
            "STATUS",
        ]
    )

    for metric, value, expected, status in qa_rows:
        qa.append(
            [
                metric,
                value,
                expected,
                status,
            ]
        )

    style_header(qa, 4)

    qa.freeze_panes = "A2"

    qa.column_dimensions["A"].width = 28
    qa.column_dimensions["B"].width = 24
    qa.column_dimensions["C"].width = 24
    qa.column_dimensions["D"].width = 14

    for row in qa.iter_rows(
        min_row=2,
        min_col=2,
        max_col=2,
    ):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000000000"

    wb.save(output_path)


def main():
    args = parse_args()

    input_path = Path(args.input).expanduser().resolve()

    if args.output:
        output_path = Path(args.output).expanduser().resolve()
    else:
        output_path = input_path.parent / OUTPUT_FILENAME

    print("=" * 86)
    print("GRAVITY v0 — E0-A TERRITORIAL INPUT MATERIALIZATION")
    print("=" * 86)

    print(f"SOURCE : {input_path}")
    print(f"SHEET  : {SOURCE_SHEET}")
    print(f"OUTPUT : {output_path}")
    print()

    if not input_path.exists():
        raise FileNotFoundError(
            f"Workbook RAW non trovato:\n{input_path}"
        )

    if input_path == output_path:
        raise ValueError(
            "INPUT e OUTPUT coincidono. "
            "Il workbook RAW non può essere sovrascritto."
        )

    if input_path.name != "Gravity_v0_territorial_inputs_raw.xlsx":
        print(
            "WARNING: il filename sorgente non coincide con "
            "Gravity_v0_territorial_inputs_raw.xlsx"
        )

    # Il RAW viene aperto esplicitamente in sola lettura.
    wb_raw = load_workbook(
        input_path,
        read_only=True,
        data_only=True,
    )

    try:
        if SOURCE_SHEET not in wb_raw.sheetnames:
            raise KeyError(
                f"Sheet richiesto '{SOURCE_SHEET}' non trovato. "
                f"Sheet disponibili: {wb_raw.sheetnames}"
            )

        ws = wb_raw[SOURCE_SHEET]

        headers = [
            cell.value
            for cell in next(ws.iter_rows(min_row=1, max_row=1))
        ]

        header_names = [
            str(value).strip()
            if value is not None
            else ""
            for value in headers
        ]

        duplicate_headers = [
            name
            for name, count in Counter(header_names).items()
            if name and count > 1
        ]

        if duplicate_headers:
            raise ValueError(
                "Header duplicati nel RAW: "
                + ", ".join(duplicate_headers)
            )

        missing_headers = [
            col
            for col in REQUIRED_COLUMNS
            if col not in header_names
        ]

        if missing_headers:
            raise KeyError(
                "Campi richiesti mancanti: "
                + ", ".join(missing_headers)
            )

        col_index = {
            name: header_names.index(name)
            for name in REQUIRED_COLUMNS
        }

        records = []

        missing_values = 0
        non_numeric_values = 0
        negative_values = 0

        for excel_row_number, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            values = {
                col: (
                    row[col_index[col]]
                    if col_index[col] < len(row)
                    else None
                )
                for col in REQUIRED_COLUMNS
            }

            # Ignora esclusivamente righe completamente vuote.
            if all(
                is_missing(values[col])
                for col in REQUIRED_COLUMNS
            ):
                continue

            row_missing = sum(
                is_missing(values[col])
                for col in REQUIRED_COLUMNS
            )

            missing_values += row_missing

            numeric = {}

            for col in NUMERIC_RAW_COLUMNS:
                if is_missing(values[col]):
                    numeric[col] = None
                    continue

                number = numeric_value(values[col])

                if number is None:
                    non_numeric_values += 1
                    numeric[col] = None
                    continue

                if number < 0:
                    negative_values += 1

                numeric[col] = number

            records.append(
                {
                    "SOURCE_ROW": excel_row_number,
                    "PRO_COM": values["PRO_COM"],
                    "COMUNE": values["COMUNE"],
                    "PARCO_AUTO_RAW": values["PARCO_AUTO_RAW"],
                    "TURISMO_RAW": values["TURISMO_RAW"],
                    "GDO_RAW": values["GDO_RAW"],
                    "_PARCO": numeric["PARCO_AUTO_RAW"],
                    "_TURISMO": numeric["TURISMO_RAW"],
                    "_GDO": numeric["GDO_RAW"],
                }
            )

    finally:
        wb_raw.close()

    rows = len(records)

    pro_com_values = [
        record["PRO_COM"]
        for record in records
        if not is_missing(record["PRO_COM"])
    ]

    distinct_pro_com = len(set(pro_com_values))

    print("RAW QA PRELIMINARE")
    print("-" * 86)
    print(f"ROWS                = {rows}")
    print(f"DISTINCT_PRO_COM    = {distinct_pro_com}")
    print(f"MISSING_VALUES      = {missing_values}")
    print(f"NON_NUMERIC_VALUES  = {non_numeric_values}")
    print(f"NEGATIVE_VALUES     = {negative_values}")

    structural_ok = (
        rows == EXPECTED_ROWS
        and distinct_pro_com == EXPECTED_ROWS
        and missing_values == 0
        and non_numeric_values == 0
        and negative_values == 0
    )

    if not structural_ok:
        print()
        print("STATUS = FAIL_RAW_QA")
        print(
            "Nessun workbook derivato scritto: "
            "il RAW richiede review."
        )
        return 2

    sum_parco = math.fsum(
        record["_PARCO"]
        for record in records
    )

    sum_turismo = math.fsum(
        record["_TURISMO"]
        for record in records
    )

    sum_gdo = math.fsum(
        record["_GDO"]
        for record in records
    )

    if sum_parco <= 0:
        raise ValueError(
            "SUM_PARCO_AUTO_RAW deve essere > 0."
        )

    if sum_turismo <= 0:
        raise ValueError(
            "SUM_TURISMO_RAW deve essere > 0."
        )

    if sum_gdo <= 0:
        raise ValueError(
            "SUM_GDO_RAW deve essere > 0."
        )

    # ==============================================================
    # NORMALIZZAZIONE L1
    # ==============================================================

    for record in records:
        p_i = record["_PARCO"] / sum_parco

        turismo_norm = (
            record["_TURISMO"] / sum_turismo
        )

        gdo_norm = (
            record["_GDO"] / sum_gdo
        )

        # Mix 50/50 autorizzato.
        # NON viene effettuata alcuna rinormalizzazione successiva.
        a_j = (
            0.5 * turismo_norm
            + 0.5 * gdo_norm
        )

        record["P_i"] = p_i
        record["TURISMO_NORM"] = turismo_norm
        record["GDO_NORM"] = gdo_norm
        record["A_j"] = a_j

    sum_p_i = math.fsum(
        record["P_i"]
        for record in records
    )

    sum_turismo_norm = math.fsum(
        record["TURISMO_NORM"]
        for record in records
    )

    sum_gdo_norm = math.fsum(
        record["GDO_NORM"]
        for record in records
    )

    sum_a_j = math.fsum(
        record["A_j"]
        for record in records
    )

    qa_rows = [
        (
            "ROWS",
            rows,
            EXPECTED_ROWS,
            status_exact(rows, EXPECTED_ROWS),
        ),
        (
            "DISTINCT_PRO_COM",
            distinct_pro_com,
            EXPECTED_ROWS,
            status_exact(
                distinct_pro_com,
                EXPECTED_ROWS,
            ),
        ),
        (
            "MISSING_VALUES",
            missing_values,
            0,
            status_exact(missing_values, 0),
        ),
        (
            "NON_NUMERIC_VALUES",
            non_numeric_values,
            0,
            status_exact(non_numeric_values, 0),
        ),
        (
            "NEGATIVE_VALUES",
            negative_values,
            0,
            status_exact(negative_values, 0),
        ),
        (
            "SUM_PARCO_AUTO_RAW",
            sum_parco,
            "> 0",
            "PASS" if sum_parco > 0 else "FAIL",
        ),
        (
            "SUM_TURISMO_RAW",
            sum_turismo,
            "> 0",
            "PASS" if sum_turismo > 0 else "FAIL",
        ),
        (
            "SUM_GDO_RAW",
            sum_gdo,
            "> 0",
            "PASS" if sum_gdo > 0 else "FAIL",
        ),
        (
            "SUM_P_i",
            sum_p_i,
            "1 ± 1e-12",
            status_close(sum_p_i),
        ),
        (
            "SUM_TURISMO_NORM",
            sum_turismo_norm,
            "1 ± 1e-12",
            status_close(sum_turismo_norm),
        ),
        (
            "SUM_GDO_NORM",
            sum_gdo_norm,
            "1 ± 1e-12",
            status_close(sum_gdo_norm),
        ),
        (
            "SUM_A_j",
            sum_a_j,
            "1 ± 1e-12",
            status_close(sum_a_j),
        ),
    ]

    overall_status = (
        "PASS"
        if all(
            row[3] == "PASS"
            for row in qa_rows
        )
        else "FAIL"
    )

    qa_rows.append(
        (
            "OVERALL_QA",
            overall_status,
            "PASS",
            overall_status,
        )
    )

    qa_rows.append(
        (
            "FLOAT_TOLERANCE",
            TOL,
            "1e-12",
            "INFO",
        )
    )

    # ==============================================================
    # OUTPUT WORKBOOK
    # ==============================================================

    create_output_workbook(
        output_path,
        records,
        qa_rows,
    )

    # ==============================================================
    # CONSOLE QA
    # ==============================================================

    print()
    print("QA SINTETICO")
    print("=" * 86)

    for metric, value, expected, status in qa_rows:
        print(
            f"{metric:<24} = "
            f"{fmt_number(value):<20} "
            f"EXPECTED={str(expected):<14} "
            f"STATUS={status}"
        )

    print_rank(
        "TOP 10 COMUNI PER P_i",
        records,
        "P_i",
        reverse=True,
    )

    print_rank(
        "TOP 10 COMUNI PER A_j",
        records,
        "A_j",
        reverse=True,
    )

    print_rank(
        "BOTTOM 10 COMUNI PER P_i",
        records,
        "P_i",
        reverse=False,
    )

    print_rank(
        "BOTTOM 10 COMUNI PER A_j",
        records,
        "A_j",
        reverse=False,
    )

    zero_a = [
        record
        for record in records
        if math.isclose(
            record["A_j"],
            0.0,
            rel_tol=0.0,
            abs_tol=1e-15,
        )
    ]

    print()
    print("COMUNI CON A_j = 0")
    print("-" * 86)

    if not zero_a:
        print("NONE")
    else:
        print(
            f"{'PRO_COM':>8}  "
            f"{'COMUNE':<35}  "
            f"{'TURISMO_RAW':>18}  "
            f"{'GDO_RAW':>18}"
        )

        print("-" * 86)

        for record in sorted(
            zero_a,
            key=lambda r: str(
                r["COMUNE"]
            ).casefold(),
        ):
            print(
                f"{str(record['PRO_COM']):>8}  "
                f"{str(record['COMUNE'])[:35]:<35}  "
                f"{fmt_number(record['_TURISMO']):>18}  "
                f"{fmt_number(record['_GDO']):>18}"
            )

    print()
    print(f"OUTPUT_WRITTEN = {output_path}")
    print(f"OVERALL_QA     = {overall_status}")

    if overall_status != "PASS":
        print(
            "ATTENZIONE: workbook creato ma QA non PASS. "
            "NON trasferire sul PC aziendale."
        )

    return 0 if overall_status == "PASS" else 3


if __name__ == "__main__":
    exit_code = 1

    try:
        exit_code = main()

    except Exception as exc:
        print()
        print("=" * 86)
        print("ERRORE")
        print("=" * 86)
        print(f"{type(exc).__name__}: {exc}")
        exit_code = 1

    finally:
        print("=== RUN COMPLETATA ===")

    sys.exit(exit_code)