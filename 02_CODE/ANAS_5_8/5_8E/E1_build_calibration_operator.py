"""
E1 — CONSTRUCTION OF CALIBRATION OPERATOR

PURPOSE
-------
Costruire e testare l'operatore deterministico che trasforma una domanda
OD comunale x_ij nei flussi modellati sulle tre sezioni ANAS PRIMARY:

    y_hat_a = sum_ij p_ij^a * x_ij

La quota p_ij^a viene costruita esclusivamente dal path system OSM frozen:
OD comunale -> 9 access-pair path -> PRODUCT-LAMBDA -> directed OSM edges
-> ANAS PRIMARY.

Lo script calcola inoltre:
- componente pendolare assegnata C_a;
- struttura Gravity con deterrenza esponenziale;
- un test strutturale con Q_TEST e BETA_TEST dichiarati NON calibrati;
- totale LIGHT di test.

INPUTS
------
1. OSM_OD_municipal_summary_v01.csv
2. OSM_OD_access_paths_v01.csv
3. OSM_OD_path_offsets_v01.npy
4. OSM_OD_transition_slots_v01.npy
5. osm_turn_state_edgeid_v01.npz
6. osm_directed_edges_v02.sqlite
7. Gravity_v0_territorial_inputs_derived_v01.xlsx
8. ISTAT_commuting_LIGHT_v0.xlsx
9. ANAS_5_8D_calibration_mapping_v01.csv

OUTPUTS
-------
Solo se tutti i QA passano, vengono scritti in area temporanea:

1. E1_calibration_operator_candidate_v01.npz
   Contiene il dominio OD e p_ij^a per le tre PRIMARY.

2. E1_primary_test_results_candidate_v01.csv
   Tre righe con:
   SECTION_ID
   TGMA_LIGHT_2024
   C_a_ISTAT
   N_a_TEST
   Y_hat_a_TEST

ASSUMPTIONS
-----------
- OD domain = 215 * 214 = 46010 ordered intermunicipal OD.
- Ogni OD possiede esattamente 9 access-pair path frozen.
- pair_weight = PRODUCT-LAMBDA.
- Routing impedance = TIME_B5.
- Gravity impedance = time_s_PRODUCT_LAMBDA / 60, quindi MINUTI.
- PRIMARY = 920034, 920032, 920039.
- MEASUREMENT_OPERATOR = BIDIRECTIONAL_SUM.
- Una path usa la sezione a se attraversa almeno uno dei directed edge
  frozen contenuti in RELEVANT_DIRECTED_EDGE_IDS.
- Una path viene contata una sola volta per sezione anche se incontrasse
  più di un edge rilevante della medesima sezione.
- Q_TEST e BETA_TEST sono esclusivamente valori diagnostici strutturali.
  NON sono stimati, NON sono ottimizzati e NON hanno valore metodologico.

FROZEN ARTIFACTS USED
---------------------
G_OSM_operativo v01
OD_PATH_SYSTEM_OSM v01
PRODUCT-LAMBDA path weights
TIME_B5 routing impedance
Gravity territorial inputs v01
ISTAT commuting LIGHT v0
ANAS 5.8D calibration mapping v01

FILES WRITTEN
-------------
Solo:
C:\\Tesi\\Tesi_QGIS\\03_output_temporanei\\fase_5_8E\\E1_operator\\
E1_calibration_operator_candidate_v01.npz
E1_primary_test_results_candidate_v01.csv

FILES NEVER MODIFIED
--------------------
Qualsiasi artefatto sotto:
C:\\Tesi\\Tesi_QGIS\\02_package\\

Nessun file frozen viene modificato.
Nessun file ANAS 2025 viene letto.
"""

from pathlib import Path
import csv
import math
import sqlite3

import numpy as np
from scipy.sparse import load_npz
from openpyxl import load_workbook


# =============================================================================
# CONSTANTS
# =============================================================================

ROOT = Path(r"C:\Tesi\Tesi_QGIS")

PRIMARY_ORDER = [
    "920034",
    "920032",
    "920039",
]

# SOLO TEST STRUTTURALE.
# NON calibrati, NON ottimizzati.
Q_TEST = 100_000.0
BETA_TEST_PER_MIN = 0.02

SCAN_CHUNK_SIZE = 10_000_000


# =============================================================================
# PATHS
# =============================================================================

OD_ROOT = (
    ROOT
    / "02_package"
    / "od_paths_osm_light"
)

OD_SUMMARY = (
    OD_ROOT
    / "OSM_OD_municipal_summary_v01.csv"
)

ACCESS_PATHS = (
    OD_ROOT
    / "OSM_OD_access_paths_v01.csv"
)

PATH_OFFSETS = (
    OD_ROOT
    / "OSM_OD_path_offsets_v01.npy"
)

TRANSITION_SLOTS = (
    OD_ROOT
    / "OSM_OD_transition_slots_v01.npy"
)

GRAPH_ROOT = (
    ROOT
    / "02_package"
    / "grafo_operativo_osm"
)

B5_EDGEID = (
    GRAPH_ROOT
    / "osm_turn_state_edgeid_v01.npz"
)

B2_SQLITE = (
    GRAPH_ROOT
    / "osm_directed_edges_v02.sqlite"
)

GRAVITY_INPUTS = (
    ROOT
    / "02_package"
    / "gravity_v0_inputs"
    / "Gravity_v0_territorial_inputs_derived_v01.xlsx"
)

COMMUTING = (
    ROOT
    / "02_package"
    / "gravity_v0_inputs"
    / "ISTAT_commuting_LIGHT_v0.xlsx"
)

ANAS_MAPPING = (
    ROOT
    / "02_package"
    / "anas_calibration_v0"
    / "ANAS_5_8D_calibration_mapping_v01.csv"
)

OUT_DIR = (
    ROOT
    / "03_output_temporanei"
    / "fase_5_8E"
    / "E1_operator"
)

OUT_OPERATOR = (
    OUT_DIR
    / "E1_calibration_operator_candidate_v01.npz"
)

OUT_RESULTS = (
    OUT_DIR
    / "E1_primary_test_results_candidate_v01.csv"
)


# =============================================================================
# HELPERS
# =============================================================================

def norm_procom(value):
    """
    Normalizza PRO_COM a stringa di 6 cifre.
    """
    if value is None:
        raise ValueError("PRO_COM nullo.")

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    if not text.isdigit():
        raise ValueError(
            f"PRO_COM non numerico: {value!r}"
        )

    return text.zfill(6)


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def require_file(path):
    if not path.is_file():
        raise FileNotFoundError(path)


def parse_pipe_ints(value):
    text = clean(value)

    if not text:
        return []

    return [
        int(x.strip())
        for x in text.split("|")
        if x.strip()
    ]


def parse_pipe_strings(value):
    text = clean(value)

    if not text:
        return []

    return [
        x.strip()
        for x in text.split("|")
        if x.strip()
    ]


# =============================================================================
# LOAD MUNICIPAL OD SUMMARY
# =============================================================================

def load_od_summary():
    origins = []
    destinations = []
    c_min = []

    od_index = {}

    with OD_SUMMARY.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(f)

        required = {
            "origin_PRO_COM",
            "destination_PRO_COM",
            "time_s_PRODUCT_LAMBDA",
        }

        missing = (
            required
            - set(reader.fieldnames or [])
        )

        if missing:
            raise RuntimeError(
                f"OD summary: campi mancanti {missing}"
            )

        for row in reader:

            o = norm_procom(
                row["origin_PRO_COM"]
            )

            d = norm_procom(
                row["destination_PRO_COM"]
            )

            key = (o, d)

            if key in od_index:
                raise RuntimeError(
                    f"OD duplicata: {key}"
                )

            t_s = float(
                row[
                    "time_s_PRODUCT_LAMBDA"
                ]
            )

            c = t_s / 60.0

            idx = len(origins)

            od_index[key] = idx
            origins.append(o)
            destinations.append(d)
            c_min.append(c)

    origins = np.asarray(
        origins,
        dtype="<U6",
    )

    destinations = np.asarray(
        destinations,
        dtype="<U6",
    )

    c_min = np.asarray(
        c_min,
        dtype=np.float64,
    )

    return (
        origins,
        destinations,
        c_min,
        od_index,
    )


# =============================================================================
# LOAD TERRITORIAL INPUTS
# =============================================================================

def load_territorial_inputs():
    wb = load_workbook(
        GRAVITY_INPUTS,
        read_only=True,
        data_only=True,
    )

    try:
        ws = wb["TERRITORIAL_INPUTS"]

        rows = ws.iter_rows(values_only=True)

        header = [
            clean(x)
            for x in next(rows)
        ]

        col = {
            name: i
            for i, name in enumerate(header)
        }

        required = {
            "PRO_COM",
            "P_i",
            "A_j",
        }

        missing = required - set(col)

        if missing:
            raise RuntimeError(
                f"Territorial inputs: campi mancanti {missing}"
            )

        data = {}

        for row in rows:
            procom = norm_procom(
                row[col["PRO_COM"]]
            )

            if procom in data:
                raise RuntimeError(
                    f"PRO_COM duplicato territorial: {procom}"
                )

            p_i = float(
                row[col["P_i"]]
            )

            a_j = float(
                row[col["A_j"]]
            )

            data[procom] = (
                p_i,
                a_j,
            )

    finally:
        wb.close()

    return data


# =============================================================================
# LOAD COMMUTING
# =============================================================================

def load_commuting(od_index):
    commuting = np.full(
        len(od_index),
        np.nan,
        dtype=np.float64,
    )

    seen = set()

    wb = load_workbook(
        COMMUTING,
        read_only=True,
        data_only=True,
    )

    try:
        ws = wb["COMMUTING_OD"]
        rows = ws.iter_rows(values_only=True)

        header = [
            clean(x)
            for x in next(rows)
        ]

        col = {
            name: i
            for i, name in enumerate(header)
        }

        required = {
            "ORIGIN_PRO_COM",
            "DESTINATION_PRO_COM",
            "C_ij_ISTAT_VEH_DAY",
        }

        missing = required - set(col)

        if missing:
            raise RuntimeError(
                f"Commuting: campi mancanti {missing}"
            )

        for row in rows:
            o = norm_procom(
                row[col["ORIGIN_PRO_COM"]]
            )

            d = norm_procom(
                row[col["DESTINATION_PRO_COM"]]
            )

            key = (o, d)

            if key in seen:
                raise RuntimeError(
                    f"Commuting OD duplicata: {key}"
                )

            seen.add(key)

            if key not in od_index:
                raise RuntimeError(
                    f"Commuting OD fuori dominio: {key}"
                )

            commuting[
                od_index[key]
            ] = float(
                row[
                    col[
                        "C_ij_ISTAT_VEH_DAY"
                    ]
                ]
            )

    finally:
        wb.close()

    return commuting, seen


# =============================================================================
# LOAD ANAS PRIMARY
# =============================================================================

def load_anas_primary():
    primary = {}

    with ANAS_MAPPING.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(
            f,
            delimiter=";",
        )

        for row in reader:

            section_id = clean(
                row.get("SECTION_ID")
            )

            if section_id not in PRIMARY_ORDER:
                continue

            if clean(row.get("ROLE")) != "PRIMARY":
                raise RuntimeError(
                    f"{section_id}: ROLE != PRIMARY"
                )

            if (
                clean(
                    row.get(
                        "MEASUREMENT_OPERATOR"
                    )
                )
                != "BIDIRECTIONAL_SUM"
            ):
                raise RuntimeError(
                    f"{section_id}: "
                    "measurement operator inatteso."
                )

            if (
                clean(
                    row.get(
                        "QGIS_REVIEW_STATUS"
                    )
                )
                != "CLOSED"
            ):
                raise RuntimeError(
                    f"{section_id}: "
                    "QGIS_REVIEW_STATUS != CLOSED"
                )

            primary[section_id] = {
                "observed_2024": float(
                    row["TGMA_LIGHT_2024"]
                ),
                "edge_ids": parse_pipe_ints(
                    row[
                        "RELEVANT_DIRECTED_EDGE_IDS"
                    ]
                ),
                "segment_uids": parse_pipe_strings(
                    row[
                        "OSM_PHYSICAL_SEGMENT"
                    ]
                ),
            }

    if set(primary) != set(PRIMARY_ORDER):
        raise RuntimeError(
            "Le tre PRIMARY non sono presenti "
            "esattamente nel mapping."
        )

    return primary


# =============================================================================
# LOAD ACCESS PATHS
# =============================================================================

def load_access_paths(
    od_index,
):
    n_expected = 414_090

    path_od_idx = np.empty(
        n_expected,
        dtype=np.int32,
    )

    pair_weight = np.empty(
        n_expected,
        dtype=np.float64,
    )

    sequence_start = np.empty(
        n_expected,
        dtype=np.int64,
    )

    sequence_end = np.empty(
        n_expected,
        dtype=np.int64,
    )

    path_count_by_od = np.zeros(
        len(od_index),
        dtype=np.int16,
    )

    weight_sum_by_od = np.zeros(
        len(od_index),
        dtype=np.float64,
    )

    with ACCESS_PATHS.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(f)

        required = {
            "path_idx",
            "origin_PRO_COM",
            "destination_PRO_COM",
            "pair_weight",
            "status",
            "sequence_start",
            "sequence_end",
            "canonical_route_impedance",
        }

        missing = (
            required
            - set(reader.fieldnames or [])
        )

        if missing:
            raise RuntimeError(
                f"Access paths: campi mancanti {missing}"
            )

        n = 0

        for row in reader:
            idx = int(row["path_idx"])

            if idx != n:
                raise RuntimeError(
                    f"path_idx non contiguo: "
                    f"atteso={n}, trovato={idx}"
                )

            if idx >= n_expected:
                raise RuntimeError(
                    "Più path del previsto."
                )

            if clean(row["status"]) != "FINITE_RECONSTRUCTED":
                raise RuntimeError(
                    f"path_idx={idx}: status non finite."
                )

            if (
                clean(
                    row[
                        "canonical_route_impedance"
                    ]
                )
                != "TIME_B5"
            ):
                raise RuntimeError(
                    f"path_idx={idx}: "
                    "impedenza non TIME_B5."
                )

            key = (
                norm_procom(
                    row["origin_PRO_COM"]
                ),
                norm_procom(
                    row[
                        "destination_PRO_COM"
                    ]
                ),
            )

            if key not in od_index:
                raise RuntimeError(
                    f"Access path OD fuori dominio: {key}"
                )

            od_idx = od_index[key]

            w = float(
                row["pair_weight"]
            )

            if (
                not math.isfinite(w)
                or w < 0
            ):
                raise RuntimeError(
                    f"path_idx={idx}: "
                    f"pair_weight invalido {w}"
                )

            path_od_idx[idx] = od_idx
            pair_weight[idx] = w

            sequence_start[idx] = int(
                row["sequence_start"]
            )

            sequence_end[idx] = int(
                row["sequence_end"]
            )

            path_count_by_od[
                od_idx
            ] += 1

            weight_sum_by_od[
                od_idx
            ] += w

            n += 1

    if n != n_expected:
        raise RuntimeError(
            f"Access paths: {n} righe, "
            f"attese {n_expected}."
        )

    return (
        path_od_idx,
        pair_weight,
        sequence_start,
        sequence_end,
        path_count_by_od,
        weight_sum_by_od,
    )


# =============================================================================
# B2 MAPPING QA
# =============================================================================

def check_b2_edges(primary):
    all_target_edges = []

    edge_to_section = {}

    for section_id in PRIMARY_ORDER:

        for edge_id in primary[
            section_id
        ]["edge_ids"]:

            if edge_id in edge_to_section:
                raise RuntimeError(
                    f"Edge {edge_id} usato da più PRIMARY."
                )

            edge_to_section[
                edge_id
            ] = section_id

            all_target_edges.append(
                edge_id
            )

    placeholders = ",".join(
        "?"
        for _ in all_target_edges
    )

    con = sqlite3.connect(
        f"file:{B2_SQLITE.as_posix()}?mode=ro",
        uri=True,
    )

    try:
        rows = con.execute(
            f"""
            SELECT edge_id, segment_uid
            FROM directed_edges
            WHERE edge_id IN ({placeholders})
            """,
            all_target_edges,
        ).fetchall()

    finally:
        con.close()

    found = {
        int(edge_id): clean(segment_uid)
        for edge_id, segment_uid in rows
    }

    if set(found) != set(all_target_edges):
        missing = (
            set(all_target_edges)
            - set(found)
        )

        raise RuntimeError(
            f"Target edge mancanti in B2: "
            f"{sorted(missing)}"
        )

    for section_id in PRIMARY_ORDER:
        valid_segments = set(
            primary[
                section_id
            ]["segment_uids"]
        )

        for edge_id in primary[
            section_id
        ]["edge_ids"]:

            actual_segment = found[
                edge_id
            ]

            if (
                actual_segment
                not in valid_segments
            ):
                raise RuntimeError(
                    f"{section_id}: edge {edge_id} "
                    f"appartiene a {actual_segment}, "
                    f"non a {valid_segments}."
                )

    return edge_to_section, found


# =============================================================================
# BUILD PATH -> SECTION HIT MATRIX
# =============================================================================

def build_path_hits(
    primary,
    offsets,
    transition_slots,
    b5_edgeid,
):
    n_paths = len(offsets) - 1

    # Codice:
    # 0 = nessuna PRIMARY
    # 1 = 920034
    # 2 = 920032
    # 3 = 920039
    slot_section = np.zeros(
        b5_edgeid.data.shape[0],
        dtype=np.uint8,
    )

    b5_counts = {}

    for code, section_id in enumerate(
        PRIMARY_ORDER,
        start=1,
    ):

        section_total = 0

        for edge_id in primary[
            section_id
        ]["edge_ids"]:

            mask = (
                b5_edgeid.data
                == edge_id
            )

            count = int(
                np.count_nonzero(mask)
            )

            b5_counts[
                (section_id, edge_id)
            ] = count

            if count == 0:
                raise RuntimeError(
                    f"{section_id}: edge_id "
                    f"{edge_id} assente nelle "
                    "transizioni B5."
                )

            if np.any(
                slot_section[mask] != 0
            ):
                raise RuntimeError(
                    "Collisione tra PRIMARY "
                    "sulle transizioni B5."
                )

            slot_section[
                mask
            ] = code

            section_total += count

        if section_total == 0:
            raise RuntimeError(
                f"{section_id}: nessuna "
                "transizione B5 trovata."
            )

    path_hits = np.zeros(
        (
            len(PRIMARY_ORDER),
            n_paths,
        ),
        dtype=bool,
    )

    n_slots = len(
        transition_slots
    )

    print()
    print("H. TRANSITION SLOT SCAN")
    print("-" * 100)

    print(
        "TRANSITION_SLOTS =",
        f"{n_slots:,}",
    )

    print(
        "SCAN_CHUNK_SIZE  =",
        f"{SCAN_CHUNK_SIZE:,}",
    )

    next_progress = 10

    for lo in range(
        0,
        n_slots,
        SCAN_CHUNK_SIZE,
    ):
        hi = min(
            lo + SCAN_CHUNK_SIZE,
            n_slots,
        )

        slots_chunk = (
            transition_slots[
                lo:hi
            ]
        )

        codes = (
            slot_section[
                slots_chunk
            ]
        )

        local_positions = (
            np.flatnonzero(
                codes
            )
        )

        if local_positions.size:

            global_positions = (
                local_positions
                + lo
            )

            path_idx = (
                np.searchsorted(
                    offsets,
                    global_positions,
                    side="right",
                )
                - 1
            )

            matched_codes = (
                codes[
                    local_positions
                ]
            )

            for code in range(
                1,
                len(PRIMARY_ORDER) + 1,
            ):

                section_paths = path_idx[
                    matched_codes
                    == code
                ]

                if section_paths.size:
                    path_hits[
                        code - 1,
                        np.unique(
                            section_paths
                        ),
                    ] = True

        progress = int(
            hi * 100 / n_slots
        )

        if (
            progress >= next_progress
            or hi == n_slots
        ):
            print(
                f"SCAN_PROGRESS = "
                f"{progress:3d}%"
            )

            while (
                next_progress
                <= progress
            ):
                next_progress += 10

    return (
        path_hits,
        b5_counts,
    )


# =============================================================================
# MAIN
# =============================================================================

def main():
    print("=" * 100)
    print("E1 — CONSTRUCTION OF CALIBRATION OPERATOR")
    print("=" * 100)

    # -----------------------------------------------------------------
    # A. File existence
    # -----------------------------------------------------------------

    print()
    print("A. FILE PREFLIGHT")
    print("-" * 100)

    inputs = [
        OD_SUMMARY,
        ACCESS_PATHS,
        PATH_OFFSETS,
        TRANSITION_SLOTS,
        B5_EDGEID,
        B2_SQLITE,
        GRAVITY_INPUTS,
        COMMUTING,
        ANAS_MAPPING,
    ]

    for path in inputs:
        require_file(path)
        print(
            "PASS =",
            path,
        )

    if OUT_OPERATOR.exists():
        raise FileExistsError(
            f"NO OVERWRITE: {OUT_OPERATOR}"
        )

    if OUT_RESULTS.exists():
        raise FileExistsError(
            f"NO OVERWRITE: {OUT_RESULTS}"
        )

    # -----------------------------------------------------------------
    # B. OD domain and impedance
    # -----------------------------------------------------------------

    print()
    print("B. MUNICIPAL OD DOMAIN / IMPEDANCE")
    print("-" * 100)

    (
        origins,
        destinations,
        c_min,
        od_index,
    ) = load_od_summary()

    n_od = len(origins)

    print("OD_ROWS      =", n_od)
    print(
        "C_MIN_FINITE =",
        int(
            np.count_nonzero(
                np.isfinite(c_min)
            )
        ),
        "/",
        n_od,
    )

    print(
        "C_MIN_GT_0   =",
        int(
            np.count_nonzero(
                c_min > 0
            )
        ),
        "/",
        n_od,
    )

    if n_od != 46_010:
        raise RuntimeError(
            f"OD_ROWS={n_od}, atteso 46010."
        )

    if not np.all(
        np.isfinite(c_min)
    ):
        raise RuntimeError(
            "c_ij_min non finite."
        )

    if not np.all(
        c_min > 0
    ):
        raise RuntimeError(
            "c_ij_min <= 0."
        )

    # -----------------------------------------------------------------
    # C. Territorial joins
    # -----------------------------------------------------------------

    print()
    print("C. TERRITORIAL P_i / A_j JOIN")
    print("-" * 100)

    territorial = (
        load_territorial_inputs()
    )

    p_origin = np.empty(
        n_od,
        dtype=np.float64,
    )

    a_destination = np.empty(
        n_od,
        dtype=np.float64,
    )

    territorial_join = 0

    for idx in range(n_od):

        o = origins[idx]
        d = destinations[idx]

        if (
            o not in territorial
            or d not in territorial
        ):
            continue

        p_origin[idx] = (
            territorial[o][0]
        )

        a_destination[idx] = (
            territorial[d][1]
        )

        territorial_join += 1

    print(
        "P_i / A_j join =",
        territorial_join,
        "/",
        n_od,
    )

    if territorial_join != n_od:
        raise RuntimeError(
            "Territorial join incompleto."
        )

    if not np.all(
        np.isfinite(p_origin)
    ):
        raise RuntimeError(
            "P_i non finite."
        )

    if not np.all(
        np.isfinite(a_destination)
    ):
        raise RuntimeError(
            "A_j non finite."
        )

    # -----------------------------------------------------------------
    # D. Commuting
    # -----------------------------------------------------------------

    print()
    print("D. ISTAT COMMUTING JOIN")
    print("-" * 100)

    commuting, commuting_seen = (
        load_commuting(
            od_index
        )
    )

    commuting_join = int(
        np.count_nonzero(
            np.isfinite(
                commuting
            )
        )
    )

    print(
        "COMMUTING_JOIN =",
        commuting_join,
        "/",
        n_od,
    )

    if commuting_join != n_od:
        raise RuntimeError(
            "Commuting join incompleto."
        )

    if len(commuting_seen) != n_od:
        raise RuntimeError(
            "Commuting domain incompleto."
        )

    # -----------------------------------------------------------------
    # E. ANAS
    # -----------------------------------------------------------------

    print()
    print("E. ANAS PRIMARY")
    print("-" * 100)

    primary = (
        load_anas_primary()
    )

    for section_id in PRIMARY_ORDER:

        info = primary[
            section_id
        ]

        print(
            f"{section_id} | "
            f"TGMA_2024="
            f"{info['observed_2024']:.0f} | "
            f"EDGES="
            f"{info['edge_ids']}"
        )

    (
        edge_to_section,
        b2_segments,
    ) = check_b2_edges(
        primary
    )

    print(
        "TARGET_EDGE_IDS_FOUND_B2 =",
        len(b2_segments),
        "/",
        len(edge_to_section),
    )

    # -----------------------------------------------------------------
    # F. Access paths
    # -----------------------------------------------------------------

    print()
    print("F. ACCESS-PAIR PATHS")
    print("-" * 100)

    (
        path_od_idx,
        pair_weight,
        sequence_start,
        sequence_end,
        path_count_by_od,
        weight_sum_by_od,
    ) = load_access_paths(
        od_index
    )

    n_paths = len(
        path_od_idx
    )

    print(
        "PATH_ROWS =",
        n_paths,
    )

    print(
        "PATHS_PER_OD_MIN =",
        int(
            path_count_by_od.min()
        ),
    )

    print(
        "PATHS_PER_OD_MAX =",
        int(
            path_count_by_od.max()
        ),
    )

    weight_error = float(
        np.max(
            np.abs(
                weight_sum_by_od
                - 1.0
            )
        )
    )

    print(
        "MAX_OD_PAIR_WEIGHT_SUM_ERROR =",
        f"{weight_error:.3e}",
    )

    if not np.all(
        path_count_by_od == 9
    ):
        raise RuntimeError(
            "Non tutte le OD hanno 9 path."
        )

    if weight_error > 1e-10:
        raise RuntimeError(
            "PRODUCT-LAMBDA weights "
            "non sommano a 1."
        )

    # -----------------------------------------------------------------
    # G. Sequence files
    # -----------------------------------------------------------------

    print()
    print("G. FROZEN SEQUENCE SYSTEM")
    print("-" * 100)

    offsets = np.load(
        PATH_OFFSETS,
        mmap_mode="r",
    )

    transition_slots = np.load(
        TRANSITION_SLOTS,
        mmap_mode="r",
    )

    b5_edgeid = (
        load_npz(
            B5_EDGEID
        )
        .tocsr()
    )

    print(
        "OFFSETS =",
        len(offsets),
    )

    print(
        "TRANSITION_SLOTS =",
        len(transition_slots),
    )

    print(
        "B5_EDGEID_NNZ =",
        b5_edgeid.nnz,
    )

    if (
        len(offsets)
        != n_paths + 1
    ):
        raise RuntimeError(
            "Offsets cardinality mismatch."
        )

    if (
        int(offsets[0])
        != 0
    ):
        raise RuntimeError(
            "offsets[0] != 0."
        )

    if (
        int(offsets[-1])
        != len(
            transition_slots
        )
    ):
        raise RuntimeError(
            "offsets[-1] != "
            "transition slot count."
        )

    if not np.array_equal(
        sequence_start,
        offsets[:-1],
    ):
        raise RuntimeError(
            "sequence_start != offsets[:-1]."
        )

    if not np.array_equal(
        sequence_end,
        offsets[1:],
    ):
        raise RuntimeError(
            "sequence_end != offsets[1:]."
        )

    print(
        "ACCESS_PATH_SEQUENCE_JOIN = PASS"
    )

    # -----------------------------------------------------------------
    # H. Scan and build path hits
    # -----------------------------------------------------------------

    (
        path_hits,
        b5_counts,
    ) = build_path_hits(
        primary,
        offsets,
        transition_slots,
        b5_edgeid,
    )

    print()
    print("I. TARGET EDGE QA")
    print("-" * 100)

    for section_id in PRIMARY_ORDER:

        for edge_id in primary[
            section_id
        ]["edge_ids"]:

            print(
                f"{section_id} | "
                f"EDGE_ID={edge_id} | "
                f"B5_SLOT_OCCURRENCES="
                f"{b5_counts[(section_id, edge_id)]}"
            )

    # -----------------------------------------------------------------
    # J. Build p_ij^a
    # -----------------------------------------------------------------

    print()
    print("J. BUILD p_ij^a")
    print("-" * 100)

    p_operator = np.zeros(
        (
            n_od,
            len(PRIMARY_ORDER),
        ),
        dtype=np.float64,
    )

    for s in range(
        len(PRIMARY_ORDER)
    ):

        hit_mask = (
            path_hits[s]
        )

        np.add.at(
            p_operator[:, s],
            path_od_idx[
                hit_mask
            ],
            pair_weight[
                hit_mask
            ],
        )

        section_id = (
            PRIMARY_ORDER[s]
        )

        print(
            f"{section_id} | "
            f"PATH_HITS="
            f"{int(np.count_nonzero(hit_mask))} | "
            f"OD_P_GT_0="
            f"{int(np.count_nonzero(p_operator[:, s] > 0))} | "
            f"P_MAX="
            f"{float(p_operator[:, s].max()):.12f}"
        )

    if not np.all(
        np.isfinite(
            p_operator
        )
    ):
        raise RuntimeError(
            "p_ij^a non finite."
        )

    if np.any(
        p_operator < -1e-12
    ):
        raise RuntimeError(
            "p_ij^a negativi."
        )

    if np.any(
        p_operator > 1.0 + 1e-10
    ):
        raise RuntimeError(
            "p_ij^a > 1."
        )

    # -----------------------------------------------------------------
    # K. Commuting assignment
    # -----------------------------------------------------------------

    print()
    print("K. COMMUTING ASSIGNMENT")
    print("-" * 100)

    C_a = (
        p_operator.T
        @ commuting
    )

    if not np.all(
        np.isfinite(C_a)
    ):
        raise RuntimeError(
            "C_a non finite."
        )

    print(
        "C_a finite =",
        int(
            np.count_nonzero(
                np.isfinite(C_a)
            )
        ),
        "/3",
    )

    # -----------------------------------------------------------------
    # L. Gravity structural test
    # -----------------------------------------------------------------

    print()
    print("L. GRAVITY STRUCTURAL TEST")
    print("-" * 100)

    print(
        "Q_TEST =",
        Q_TEST,
        "VEH/DAY",
    )

    print(
        "BETA_TEST =",
        BETA_TEST_PER_MIN,
        "PER_MIN",
    )

    print(
        "TEST_VALUES_STATUS = "
        "STRUCTURAL_ONLY / NOT_CALIBRATED"
    )

    W = (
        p_origin
        * a_destination
        * np.exp(
            -BETA_TEST_PER_MIN
            * c_min
        )
    )

    if not np.all(
        np.isfinite(W)
    ):
        raise RuntimeError(
            "W_ij non finite."
        )

    sum_W = float(
        W.sum()
    )

    if (
        not math.isfinite(sum_W)
        or sum_W <= 0
    ):
        raise RuntimeError(
            "SUM W invalida."
        )

    S = W / sum_W

    N = Q_TEST * S

    sum_S = float(
        S.sum()
    )

    sum_N = float(
        N.sum()
    )

    print(
        "SUM_S_ij =",
        f"{sum_S:.15f}",
    )

    print(
        "SUM_N_ij =",
        f"{sum_N:.12f}",
    )

    print(
        "ABS_SUM_N_MINUS_Q =",
        f"{abs(sum_N - Q_TEST):.3e}",
    )

    if abs(
        sum_S - 1.0
    ) > 1e-12:
        raise RuntimeError(
            "SUM S_ij != 1."
        )

    if abs(
        sum_N - Q_TEST
    ) > 1e-8:
        raise RuntimeError(
            "SUM N_ij != Q_TEST."
        )

    N_a = (
        p_operator.T
        @ N
    )

    Y_hat = (
        C_a
        + N_a
    )

    if not np.all(
        np.isfinite(N_a)
    ):
        raise RuntimeError(
            "N_a non finite."
        )

    if not np.all(
        np.isfinite(Y_hat)
    ):
        raise RuntimeError(
            "Y_hat non finite."
        )

    print(
        "N_a finite =",
        int(
            np.count_nonzero(
                np.isfinite(N_a)
            )
        ),
        "/3",
    )

    print(
        "Y_hat finite =",
        int(
            np.count_nonzero(
                np.isfinite(Y_hat)
            )
        ),
        "/3",
    )

    # -----------------------------------------------------------------
    # M. Mandatory QA
    # -----------------------------------------------------------------

    print()
    print("M. E1 MANDATORY QA")
    print("=" * 100)

    qa = {
        "OD_ROWS_46010":
            n_od == 46_010,

        "P_i_A_j_JOIN_46010":
            territorial_join
            == 46_010,

        "C_IJ_JOIN_46010":
            commuting_join
            == 46_010,

        "C_IJ_MIN_FINITE_46010":
            bool(
                np.all(
                    np.isfinite(
                        c_min
                    )
                )
            ),

        "C_IJ_MIN_GT_0_46010":
            bool(
                np.all(
                    c_min > 0
                )
            ),

        "NINE_PATHS_PER_OD":
            bool(
                np.all(
                    path_count_by_od
                    == 9
                )
            ),

        "PRODUCT_LAMBDA_SUMS_1":
            weight_error
            <= 1e-10,

        "NO_SEQUENCE_MISSING":
            (
                len(offsets)
                == n_paths + 1
                and int(
                    offsets[-1]
                )
                == len(
                    transition_slots
                )
            ),

        "SUM_S_EQ_1":
            abs(
                sum_S - 1.0
            )
            <= 1e-12,

        "SUM_N_EQ_Q_TEST":
            abs(
                sum_N - Q_TEST
            )
            <= 1e-8,

        "C_A_FINITE_3_OF_3":
            bool(
                np.all(
                    np.isfinite(
                        C_a
                    )
                )
            ),

        "N_A_FINITE_3_OF_3":
            bool(
                np.all(
                    np.isfinite(
                        N_a
                    )
                )
            ),

        "Y_HAT_FINITE_3_OF_3":
            bool(
                np.all(
                    np.isfinite(
                        Y_hat
                    )
                )
            ),
    }

    for name, passed in qa.items():
        print(
            f"{name:<30} = "
            f"{'PASS' if passed else 'FAIL'}"
        )

    if not all(
        qa.values()
    ):
        raise RuntimeError(
            "E1 mandatory QA failed."
        )

    # -----------------------------------------------------------------
    # N. Results
    # -----------------------------------------------------------------

    print()
    print("N. PRIMARY STRUCTURAL TEST RESULTS")
    print("=" * 100)

    header = (
        f"{'SECTION_ID':<12}"
        f"{'TGMA_LIGHT_2024':>18}"
        f"{'C_a_ISTAT':>18}"
        f"{'N_a_TEST':>18}"
        f"{'Y_hat_a_TEST':>18}"
    )

    print(header)
    print("-" * len(header))

    results_rows = []

    for s, section_id in enumerate(
        PRIMARY_ORDER
    ):

        observed = primary[
            section_id
        ]["observed_2024"]

        row = {
            "SECTION_ID":
                section_id,
            "TGMA_LIGHT_2024":
                observed,
            "C_a_ISTAT":
                float(C_a[s]),
            "N_a_TEST":
                float(N_a[s]),
            "Y_hat_a_TEST":
                float(Y_hat[s]),
        }

        results_rows.append(
            row
        )

        print(
            f"{section_id:<12}"
            f"{observed:>18.3f}"
            f"{C_a[s]:>18.3f}"
            f"{N_a[s]:>18.3f}"
            f"{Y_hat[s]:>18.3f}"
        )

    # -----------------------------------------------------------------
    # O. Save TEMPORARY candidate operator only after PASS
    # -----------------------------------------------------------------

    print()
    print("O. TEMPORARY E1 OUTPUT")
    print("-" * 100)

    OUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if (
        OUT_OPERATOR.exists()
        or OUT_RESULTS.exists()
    ):
        raise FileExistsError(
            "NO OVERWRITE: "
            "E1 output already exists."
        )

    np.savez_compressed(
        OUT_OPERATOR,
        origin_PRO_COM=origins,
        destination_PRO_COM=destinations,
        c_ij_min=c_min,
        P_i=p_origin,
        A_j=a_destination,
        C_ij_ISTAT=commuting,
        p_920034=p_operator[:, 0],
        p_920032=p_operator[:, 1],
        p_920039=p_operator[:, 2],
        Q_TEST=np.asarray(
            [Q_TEST],
            dtype=np.float64,
        ),
        BETA_TEST_PER_MIN=np.asarray(
            [BETA_TEST_PER_MIN],
            dtype=np.float64,
        ),
    )

    with OUT_RESULTS.open(
        "x",
        encoding="utf-8",
        newline="",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=[
                "SECTION_ID",
                "TGMA_LIGHT_2024",
                "C_a_ISTAT",
                "N_a_TEST",
                "Y_hat_a_TEST",
            ],
        )

        writer.writeheader()
        writer.writerows(
            results_rows
        )

    print(
        "OPERATOR_CANDIDATE =",
        OUT_OPERATOR,
    )

    print(
        "PRIMARY_RESULTS    =",
        OUT_RESULTS,
    )

    print(
        "OUTPUT_STATUS = "
        "TEMPORARY / NOT_FROZEN / "
        "NOT_PROMOTED"
    )

    # -----------------------------------------------------------------
    # P. Final E1 state
    # -----------------------------------------------------------------

    print()
    print("P. E1 STATE")
    print("=" * 100)

    print(
        "E1_STRUCTURAL_TEST = PASS"
    )

    print(
        "E1_OPERATOR_BUILT  = YES"
    )

    print(
        "LOSS_SELECTED      = NO"
    )

    print(
        "Q_CALIBRATED       = NO"
    )

    print(
        "BETA_CALIBRATED    = NO"
    )

    print(
        "ANAS_2025_USED     = NO"
    )

    print(
        "E2_STARTED         = NO"
    )

    print(
        "NEXT_ACTION        = "
        "RETURN E1 RESULTS TO CHAT"
    )


if __name__ == "__main__":
    try:
        main()
    finally:
        print("=== RUN COMPLETATA ===")


