"""
E0-C — MATERIALIZZAZIONE CONTROLLATA C_ij^ISTAT

PURPOSE
-------
Materializzare la componente pendolare giornaliera LIGHT v0 C_ij^ISTAT
sul dominio completo delle 46.010 OD ordinate intercomunali FVG.

Per ogni relazione osservata o -> d con Pendolari = P_od:

    contributo andata:
        C_OUTBOUND_CONTRIBUTION = P_od * ALPHA_PEND

    contributo di ritorno verso o -> d generato dalla relazione osservata d -> o:
        C_RETURN_CONTRIBUTION = P_do * ALPHA_PEND

    flusso finale:
        C_ij_ISTAT_VEH_DAY =
            C_OUTBOUND_CONTRIBUTION
            + C_RETURN_CONTRIBUTION

Quindi:

    C_ij = ALPHA_PEND * (P_ij + P_ji)

La matrice finale deve essere simmetrica.
NON viene applicato alcun ulteriore fattore 2.

INPUTS
------
C:\\Tesi\\Tesi_QGIS\\00_originali\\
flussi_mobilita_interni_FVG_2021_tempi_distanze.csv

Campi canonici:
    Procom_res
    Procom_lav
    Pendolari

Campi Comune_res e Comune_lav, se presenti, vengono utilizzati
esclusivamente per rendere leggibili gli esempi stampati in console.

OUTPUTS
-------
C:\\Tesi\\Tesi_QGIS\\03_output_temporanei\\fase_5_8E\\
ISTAT_commuting_LIGHT_v0_candidate.xlsx

Fogli:
    COMMUTING_OD
    QA

ASSUMPTIONS
-----------
1. Baseline esclusivamente FVG-only.
2. Le OD intracomunali i == j sono escluse.
3. Il dominio finale contiene tutte le OD ordinate intercomunali:
       215 * 214 = 46.010.
4. OD senza contributi pendolari ricevono C_ij_ISTAT = 0.
5. Eventuali righe raw multiple sulla stessa OD vengono aggregate.
6. Ogni relazione raw direzionale viene trattata autonomamente.
7. ALPHA_PEND è congelato al valore numerico 0.42855.
8. Gli eventuali Pendolari intracomunali presenti nel raw non vengono
   utilizzati nella matrice intercomunale.
9. Non viene effettuata alcuna calibrazione Gravity.
10. Tolleranza floating point per la simmetria = 1e-12.

FROZEN ARTIFACTS USED
---------------------
Raw ISTAT FVG 2021:
    flussi_mobilita_interni_FVG_2021_tempi_distanze.csv

Regola frozen:
    ALPHA_PEND = 0.42855
    ANDATA + RITORNO
    INTRAZONAL EXCLUDED
    FVG_ONLY

FILES WRITTEN
-------------
Solo:
    ISTAT_commuting_LIGHT_v0_candidate.xlsx

FILES NEVER MODIFIED
--------------------
flussi_mobilita_interni_FVG_2021_tempi_distanze.csv

Qualsiasi artefatto sotto:
    C:\\Tesi\\Tesi_QGIS\\02_package\\

Nessun file frozen viene modificato.
"""

from __future__ import annotations

import csv
import hashlib
import math
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ======================================================================
# CONFIGURAZIONE FROZEN
# ======================================================================

INPUT_PATH = Path(
    r"C:\Tesi\Tesi_QGIS\00_originali"
    r"\flussi_mobilita_interni_FVG_2021_tempi_distanze.csv"
)

OUTPUT_DIR = Path(
    r"C:\Tesi\Tesi_QGIS\03_output_temporanei\fase_5_8E"
)

OUTPUT_PATH = OUTPUT_DIR / "ISTAT_commuting_LIGHT_v0_candidate.xlsx"

ORIGIN_FIELD = "Procom_res"
DESTINATION_FIELD = "Procom_lav"
RAW_FLOW_FIELD = "Pendolari"

OPTIONAL_ORIGIN_NAME_FIELD = "Comune_res"
OPTIONAL_DESTINATION_NAME_FIELD = "Comune_lav"

ALPHA_PEND = 0.42855

EXPECTED_MUNICIPALITIES = 215
EXPECTED_ORDERED_OD = 215 * 214  # 46.010

SYMMETRY_TOL = 1e-12


# ======================================================================
# FUNZIONI DI SUPPORTO
# ======================================================================


def normalize_pro_com(value: str, field_name: str, row_number: int) -> str:
    """
    Normalizza PRO_COM come stringa di 6 cifre, preservando gli zeri iniziali.
    """

    if value is None:
        raise ValueError(
            f"Riga {row_number}: {field_name} mancante."
        )

    text = str(value).strip()

    if not text:
        raise ValueError(
            f"Riga {row_number}: {field_name} vuoto."
        )

    if not text.isdigit():
        raise ValueError(
            f"Riga {row_number}: {field_name} non numerico: {text!r}"
        )

    if len(text) > 6:
        raise ValueError(
            f"Riga {row_number}: {field_name} ha più di 6 cifre: {text!r}"
        )

    return text.zfill(6)


def parse_nonnegative_number(
    value: str,
    field_name: str,
    row_number: int,
) -> float:
    """
    Legge un valore numerico non negativo.
    Nessuna imputazione o correzione automatica.
    """

    if value is None or str(value).strip() == "":
        raise ValueError(
            f"Riga {row_number}: {field_name} mancante."
        )

    text = str(value).strip()

    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(
            f"Riga {row_number}: "
            f"{field_name} non numerico: {text!r}"
        ) from exc

    if not math.isfinite(number):
        raise ValueError(
            f"Riga {row_number}: "
            f"{field_name} non finito: {text!r}"
        )

    if number < 0:
        raise ValueError(
            f"Riga {row_number}: "
            f"{field_name} negativo: {number}"
        )

    return number


def display_number(value: float) -> str:
    """
    Formato compatto per console.
    """

    if math.isclose(value, round(value), abs_tol=1e-12):
        return str(int(round(value)))

    return f"{value:.6f}"


def sha256_file(path: Path) -> str:
    """
    SHA256 streaming del file.
    """

    digest = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def style_header(ws) -> None:
    """
    Styling minimale per leggibilità Excel.
    Non altera i dati.
    """

    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


# ======================================================================
# MAIN
# ======================================================================


def main() -> int:

    print("=" * 100)
    print("GRAVITY v0 — E0-C MATERIALIZZAZIONE C_ij^ISTAT")
    print("=" * 100)

    print(f"INPUT       = {INPUT_PATH}")
    print(f"OUTPUT      = {OUTPUT_PATH}")
    print(f"ALPHA_PEND  = {ALPHA_PEND}")
    print()

    # ------------------------------------------------------------------
    # PREFLIGHT
    # ------------------------------------------------------------------

    if not INPUT_PATH.exists():
        raise FileNotFoundError(
            f"Raw ISTAT non trovato:\n{INPUT_PATH}"
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if OUTPUT_PATH.exists():
        raise FileExistsError(
            "OVERWRITE VIETATO. "
            f"Il candidato esiste già:\n{OUTPUT_PATH}"
        )

    # ------------------------------------------------------------------
    # LETTURA RAW
    # ------------------------------------------------------------------

    raw_flow: dict[tuple[str, str], float] = defaultdict(float)

    municipalities: set[str] = set()

    # Solo per gli esempi console.
    municipality_names: dict[str, str] = {}

    source_rows = 0
    source_intermunicipal_rows = 0
    source_intrazonal_rows = 0

    sum_raw_pendolari_all = 0.0
    sum_raw_pendolari_used = 0.0
    sum_raw_pendolari_intrazonal_excluded = 0.0

    with INPUT_PATH.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(f)

        if reader.fieldnames is None:
            raise RuntimeError(
                "Il CSV non contiene un header."
            )

        required_fields = {
            ORIGIN_FIELD,
            DESTINATION_FIELD,
            RAW_FLOW_FIELD,
        }

        missing_fields = required_fields - set(reader.fieldnames)

        if missing_fields:
            raise KeyError(
                "Campi canonici mancanti nel raw: "
                + ", ".join(sorted(missing_fields))
            )

        has_names = (
            OPTIONAL_ORIGIN_NAME_FIELD in reader.fieldnames
            and OPTIONAL_DESTINATION_NAME_FIELD in reader.fieldnames
        )

        for row_number, row in enumerate(reader, start=2):

            source_rows += 1

            origin = normalize_pro_com(
                row[ORIGIN_FIELD],
                ORIGIN_FIELD,
                row_number,
            )

            destination = normalize_pro_com(
                row[DESTINATION_FIELD],
                DESTINATION_FIELD,
                row_number,
            )

            pendolari = parse_nonnegative_number(
                row[RAW_FLOW_FIELD],
                RAW_FLOW_FIELD,
                row_number,
            )

            municipalities.add(origin)
            municipalities.add(destination)

            sum_raw_pendolari_all += pendolari

            if has_names:

                origin_name = str(
                    row.get(
                        OPTIONAL_ORIGIN_NAME_FIELD,
                        "",
                    )
                ).strip()

                destination_name = str(
                    row.get(
                        OPTIONAL_DESTINATION_NAME_FIELD,
                        "",
                    )
                ).strip()

                if origin_name and origin not in municipality_names:
                    municipality_names[origin] = origin_name

                if (
                    destination_name
                    and destination not in municipality_names
                ):
                    municipality_names[destination] = destination_name

            # Intrazonali esclusi dal dominio finale.
            if origin == destination:

                source_intrazonal_rows += 1

                sum_raw_pendolari_intrazonal_excluded += pendolari

                continue

            source_intermunicipal_rows += 1
            sum_raw_pendolari_used += pendolari

            # Aggregazione trasparente di eventuali righe duplicate.
            raw_flow[(origin, destination)] += pendolari

    # ------------------------------------------------------------------
    # QA DEL DOMINIO
    # ------------------------------------------------------------------

    municipality_list = sorted(municipalities)

    if len(municipality_list) != EXPECTED_MUNICIPALITIES:
        raise RuntimeError(
            "Numero comuni inatteso nel raw: "
            f"{len(municipality_list)} "
            f"(atteso {EXPECTED_MUNICIPALITIES})."
        )

    # ------------------------------------------------------------------
    # COSTRUZIONE FULL OD 215 x 214
    # ------------------------------------------------------------------

    records = []

    for origin in municipality_list:

        for destination in municipality_list:

            if origin == destination:
                continue

            p_od = raw_flow.get(
                (origin, destination),
                0.0,
            )

            p_do = raw_flow.get(
                (destination, origin),
                0.0,
            )

            outbound = p_od * ALPHA_PEND

            return_contribution = p_do * ALPHA_PEND

            c_ij = outbound + return_contribution

            records.append(
                {
                    "ORIGIN_PRO_COM": origin,
                    "DESTINATION_PRO_COM": destination,
                    "PENDOLARI_OBSERVED_OD": p_od,
                    "PENDOLARI_OBSERVED_REVERSE": p_do,
                    "C_OUTBOUND_CONTRIBUTION": outbound,
                    "C_RETURN_CONTRIBUTION": return_contribution,
                    "C_ij_ISTAT_VEH_DAY": c_ij,
                }
            )

    # ------------------------------------------------------------------
    # QA FINALE
    # ------------------------------------------------------------------

    rows = len(records)

    unique_ordered_od = len(
        {
            (
                r["ORIGIN_PRO_COM"],
                r["DESTINATION_PRO_COM"],
            )
            for r in records
        }
    )

    intrazonal_output = sum(
        1
        for r in records
        if r["ORIGIN_PRO_COM"] == r["DESTINATION_PRO_COM"]
    )

    missing_output = 0
    negative_output = 0

    numeric_fields = [
        "PENDOLARI_OBSERVED_OD",
        "PENDOLARI_OBSERVED_REVERSE",
        "C_OUTBOUND_CONTRIBUTION",
        "C_RETURN_CONTRIBUTION",
        "C_ij_ISTAT_VEH_DAY",
    ]

    for record in records:

        for field in numeric_fields:

            value = record[field]

            if value is None or not math.isfinite(value):
                missing_output += 1

            elif value < 0:
                negative_output += 1

        if not record["ORIGIN_PRO_COM"]:
            missing_output += 1

        if not record["DESTINATION_PRO_COM"]:
            missing_output += 1

    n_positive = sum(
        1
        for r in records
        if r["C_ij_ISTAT_VEH_DAY"] > 0
    )

    n_zero = sum(
        1
        for r in records
        if r["C_ij_ISTAT_VEH_DAY"] == 0
    )

    sum_c_ij = math.fsum(
        r["C_ij_ISTAT_VEH_DAY"]
        for r in records
    )

    c_lookup = {
        (
            r["ORIGIN_PRO_COM"],
            r["DESTINATION_PRO_COM"],
        ): r["C_ij_ISTAT_VEH_DAY"]
        for r in records
    }

    max_abs_symmetry_difference = max(
        abs(
            value
            - c_lookup[(destination, origin)]
        )
        for (origin, destination), value
        in c_lookup.items()
    )

    qa_pass = (
        rows == EXPECTED_ORDERED_OD
        and unique_ordered_od == EXPECTED_ORDERED_OD
        and intrazonal_output == 0
        and missing_output == 0
        and negative_output == 0
        and math.isclose(
            max_abs_symmetry_difference,
            0.0,
            rel_tol=0.0,
            abs_tol=SYMMETRY_TOL,
        )
    )

    if not qa_pass:
        raise RuntimeError(
            "QA finale FAIL. "
            "Il workbook candidato NON verrà scritto."
        )

    # ------------------------------------------------------------------
    # SELEZIONE DETERMINISTICA DI 10 ESEMPI
    # ------------------------------------------------------------------

    both_directions = [
        r
        for r in records
        if (
            r["PENDOLARI_OBSERVED_OD"] > 0
            and r["PENDOLARI_OBSERVED_REVERSE"] > 0
        )
    ]

    one_direction = [
        r
        for r in records
        if (
            r["PENDOLARI_OBSERVED_OD"] > 0
            and r["PENDOLARI_OBSERVED_REVERSE"] == 0
        )
    ]

    zero_flows = [
        r
        for r in records
        if r["C_ij_ISTAT_VEH_DAY"] == 0
    ]

    if not both_directions:
        raise RuntimeError(
            "Nessuna coppia con flussi osservati in entrambi i versi."
        )

    if not one_direction:
        raise RuntimeError(
            "Nessuna coppia con un solo verso osservato."
        )

    if not zero_flows:
        raise RuntimeError(
            "Nessuna OD a flusso finale zero disponibile per la review."
        )

    # Per i casi positivi scegliamo esempi con massa maggiore:
    # sono più semplici da controllare manualmente.
    both_directions.sort(
        key=lambda r: (
            -r["C_ij_ISTAT_VEH_DAY"],
            r["ORIGIN_PRO_COM"],
            r["DESTINATION_PRO_COM"],
        )
    )

    one_direction.sort(
        key=lambda r: (
            -r["C_ij_ISTAT_VEH_DAY"],
            r["ORIGIN_PRO_COM"],
            r["DESTINATION_PRO_COM"],
        )
    )

    zero_flows.sort(
        key=lambda r: (
            r["ORIGIN_PRO_COM"],
            r["DESTINATION_PRO_COM"],
        )
    )

    examples = []

    def add_example(label: str, record: dict) -> None:

        key = (
            record["ORIGIN_PRO_COM"],
            record["DESTINATION_PRO_COM"],
        )

        existing = {
            (
                item[1]["ORIGIN_PRO_COM"],
                item[1]["DESTINATION_PRO_COM"],
            )
            for item in examples
        }

        if key not in existing:
            examples.append((label, record))

    # Tre categorie obbligatorie.
    add_example(
        "BOTH_DIRECTIONS_OBSERVED",
        both_directions[0],
    )

    add_example(
        "ONE_DIRECTION_OBSERVED",
        one_direction[0],
    )

    add_example(
        "ZERO_C",
        zero_flows[0],
    )

    # Completa fino a 10 con OD positive ad alta massa.
    positive_ranked = sorted(
        [
            r
            for r in records
            if r["C_ij_ISTAT_VEH_DAY"] > 0
        ],
        key=lambda r: (
            -r["C_ij_ISTAT_VEH_DAY"],
            r["ORIGIN_PRO_COM"],
            r["DESTINATION_PRO_COM"],
        ),
    )

    for record in positive_ranked:

        if len(examples) >= 10:
            break

        add_example(
            "ADDITIONAL_POSITIVE",
            record,
        )

    if len(examples) != 10:
        raise RuntimeError(
            f"Impossibile costruire 10 esempi: {len(examples)} trovati."
        )

    # ------------------------------------------------------------------
    # CREAZIONE WORKBOOK
    # ------------------------------------------------------------------

    wb = Workbook()

    ws = wb.active
    ws.title = "COMMUTING_OD"

    output_headers = [
        "ORIGIN_PRO_COM",
        "DESTINATION_PRO_COM",
        "PENDOLARI_OBSERVED_OD",
        "PENDOLARI_OBSERVED_REVERSE",
        "C_OUTBOUND_CONTRIBUTION",
        "C_RETURN_CONTRIBUTION",
        "C_ij_ISTAT_VEH_DAY",
    ]

    ws.append(output_headers)

    for record in records:

        ws.append(
            [
                record["ORIGIN_PRO_COM"],
                record["DESTINATION_PRO_COM"],
                record["PENDOLARI_OBSERVED_OD"],
                record["PENDOLARI_OBSERVED_REVERSE"],
                record["C_OUTBOUND_CONTRIBUTION"],
                record["C_RETURN_CONTRIBUTION"],
                record["C_ij_ISTAT_VEH_DAY"],
            ]
        )

    style_header(ws)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 23
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 27
    ws.column_dimensions["G"].width = 24

    # PRO_COM deve essere visualizzato come testo, mantenendo 6 cifre.
    for row in ws.iter_rows(
        min_row=2,
        min_col=1,
        max_col=2,
    ):
        for cell in row:
            cell.number_format = "@"

    for row in ws.iter_rows(
        min_row=2,
        min_col=5,
        max_col=7,
    ):
        for cell in row:
            cell.number_format = "0.000000000000"

    # ------------------------------------------------------------------
    # QA SHEET
    # ------------------------------------------------------------------

    qa = wb.create_sheet("QA")

    qa.append(
        [
            "METRIC",
            "VALUE",
            "EXPECTED",
            "STATUS",
        ]
    )

    qa_rows = [
        (
            "ROWS",
            rows,
            EXPECTED_ORDERED_OD,
            "PASS",
        ),
        (
            "UNIQUE_ORDERED_OD",
            unique_ordered_od,
            EXPECTED_ORDERED_OD,
            "PASS",
        ),
        (
            "INTRAZONAL",
            intrazonal_output,
            0,
            "PASS",
        ),
        (
            "MISSING",
            missing_output,
            0,
            "PASS",
        ),
        (
            "NEGATIVE",
            negative_output,
            0,
            "PASS",
        ),
        (
            "ALPHA_PEND",
            ALPHA_PEND,
            0.42855,
            "PASS",
        ),
        (
            "SOURCE_RAW_ROWS",
            source_rows,
            "INFO",
            "INFO",
        ),
        (
            "SOURCE_INTERMUNICIPAL_ROWS",
            source_intermunicipal_rows,
            "INFO",
            "INFO",
        ),
        (
            "SOURCE_INTRAZONAL_ROWS_EXCLUDED",
            source_intrazonal_rows,
            "INFO",
            "INFO",
        ),
        (
            "SUM_RAW_PENDOLARI_ALL",
            sum_raw_pendolari_all,
            "INFO",
            "INFO",
        ),
        (
            "SUM_RAW_PENDOLARI_USED",
            sum_raw_pendolari_used,
            "INFO",
            "INFO",
        ),
        (
            "SUM_RAW_PENDOLARI_INTRAZONAL_EXCLUDED",
            sum_raw_pendolari_intrazonal_excluded,
            "INFO",
            "INFO",
        ),
        (
            "SUM_C_ij_ISTAT",
            sum_c_ij,
            "INFO",
            "INFO",
        ),
        (
            "OD_C_GT_0",
            n_positive,
            "INFO",
            "INFO",
        ),
        (
            "OD_C_EQ_0",
            n_zero,
            "INFO",
            "INFO",
        ),
        (
            "MAX_ABS_SYMMETRY_DIFFERENCE",
            max_abs_symmetry_difference,
            "<= 1e-12",
            "PASS",
        ),
        (
            "OVERALL_QA",
            "PASS",
            "PASS",
            "PASS",
        ),
    ]

    for qa_row in qa_rows:
        qa.append(qa_row)

    style_header(qa)

    qa.freeze_panes = "A2"

    qa.column_dimensions["A"].width = 42
    qa.column_dimensions["B"].width = 24
    qa.column_dimensions["C"].width = 20
    qa.column_dimensions["D"].width = 14

    wb.save(OUTPUT_PATH)
    wb.close()

    # ------------------------------------------------------------------
    # READ-BACK DEL CANDIDATO
    # ------------------------------------------------------------------

    check_wb = load_workbook(
        OUTPUT_PATH,
        read_only=True,
        data_only=True,
    )

    try:

        expected_sheets = {
            "COMMUTING_OD",
            "QA",
        }

        if not expected_sheets.issubset(
            set(check_wb.sheetnames)
        ):
            raise RuntimeError(
                "Read-back: sheet richiesti mancanti."
            )

        check_ws = check_wb["COMMUTING_OD"]

        check_headers = [
            cell.value
            for cell in next(
                check_ws.iter_rows(
                    min_row=1,
                    max_row=1,
                )
            )
        ]

        if check_headers != output_headers:
            raise RuntimeError(
                "Read-back: header COMMUTING_OD inatteso."
            )

        check_rows = check_ws.max_row - 1

        if check_rows != EXPECTED_ORDERED_OD:
            raise RuntimeError(
                "Read-back: numero righe inatteso: "
                f"{check_rows}"
            )

    finally:
        check_wb.close()

    output_sha256 = sha256_file(OUTPUT_PATH)

    # ------------------------------------------------------------------
    # CONSOLE REPORT
    # ------------------------------------------------------------------

    print("SOURCE QA")
    print("-" * 100)

    print(
        f"SOURCE_RAW_ROWS                    = {source_rows}"
    )
    print(
        f"SOURCE_INTERMUNICIPAL_ROWS         = "
        f"{source_intermunicipal_rows}"
    )
    print(
        f"SOURCE_INTRAZONAL_ROWS_EXCLUDED    = "
        f"{source_intrazonal_rows}"
    )
    print(
        f"SUM_RAW_PENDOLARI_ALL              = "
        f"{display_number(sum_raw_pendolari_all)}"
    )
    print(
        f"SUM_RAW_PENDOLARI_USED             = "
        f"{display_number(sum_raw_pendolari_used)}"
    )
    print(
        f"SUM_RAW_PENDOLARI_INTRAZ_EXCLUDED  = "
        f"{display_number(sum_raw_pendolari_intrazonal_excluded)}"
    )

    print()
    print("FINAL QA")
    print("=" * 100)

    print(
        f"ROWS                               = {rows}"
    )
    print(
        f"UNIQUE_ORDERED_OD                  = "
        f"{unique_ordered_od}"
    )
    print(
        f"INTRAZONAL                         = "
        f"{intrazonal_output}"
    )
    print(
        f"MISSING                            = "
        f"{missing_output}"
    )
    print(
        f"NEGATIVE                           = "
        f"{negative_output}"
    )
    print(
        f"ALPHA_PEND                         = "
        f"{ALPHA_PEND}"
    )
    print(
        f"SUM_C_ij_ISTAT                     = "
        f"{sum_c_ij:.12f}"
    )
    print(
        f"OD_C_GT_0                          = "
        f"{n_positive}"
    )
    print(
        f"OD_C_EQ_0                          = "
        f"{n_zero}"
    )
    print(
        f"MAX_ABS_SYMMETRY_DIFFERENCE        = "
        f"{max_abs_symmetry_difference:.16g}"
    )
    print(
        "OVERALL_QA                         = PASS"
    )

    # ------------------------------------------------------------------
    # REVIEW UMANA — 10 ESEMPI
    # ------------------------------------------------------------------

    print()
    print("10 ESEMPI PER REVIEW UMANA")
    print("=" * 100)

    for index, (label, record) in enumerate(
        examples,
        start=1,
    ):

        origin = record["ORIGIN_PRO_COM"]
        destination = record["DESTINATION_PRO_COM"]

        origin_name = municipality_names.get(
            origin,
            "N/A",
        )

        destination_name = municipality_names.get(
            destination,
            "N/A",
        )

        print()
        print(
            f"[{index:02d}] {label}"
        )
        print(
            f"ORIGIN      = {origin} | {origin_name}"
        )
        print(
            f"DESTINATION = "
            f"{destination} | {destination_name}"
        )
        print(
            "Pendolari o->d               = "
            f"{display_number(record['PENDOLARI_OBSERVED_OD'])}"
        )
        print(
            "Pendolari d->o               = "
            f"{display_number(record['PENDOLARI_OBSERVED_REVERSE'])}"
        )
        print(
            "Contributo andata            = "
            f"{record['C_OUTBOUND_CONTRIBUTION']:.12f}"
        )
        print(
            "Contributo ritorno           = "
            f"{record['C_RETURN_CONTRIBUTION']:.12f}"
        )
        print(
            "C_ij finale                  = "
            f"{record['C_ij_ISTAT_VEH_DAY']:.12f}"
        )

    print()
    print("=" * 100)
    print(f"OUTPUT_WRITTEN = {OUTPUT_PATH}")
    print(f"SHA256         = {output_sha256}")
    print("FILE_READABLE  = CONFIRMED")
    print("PROMOTED       = NO")
    print("E1_STARTED     = NO")

    return 0


if __name__ == "__main__":

    exit_code = 1

    try:
        exit_code = main()

    except Exception as exc:
        print()
        print("=" * 100)
        print("ERRORE")
        print("=" * 100)
        print(f"{type(exc).__name__}: {exc}")
        exit_code = 1

    finally:
        print("=== RUN COMPLETATA ===")

    sys.exit(exit_code)